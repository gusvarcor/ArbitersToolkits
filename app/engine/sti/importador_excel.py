from openpyxl import load_workbook

from app.engine.sti.jugador import Jugador


def importar_excel(ruta_excel):
    """
    Lee un archivo de Excel y devuelve
    una lista de objetos Jugador.
    """

    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb.active

    jugadores = []

    # Empieza en la fila 2 porque la fila 1 son encabezados
    for fila in ws.iter_rows(min_row=2, values_only=True):

        jugador = Jugador()

        jugador.id = fila[0]
        jugador.grupo = fila[1]

        jugador.apellido_paterno = fila[2] or ""
        jugador.apellido_materno = fila[3] or ""
        jugador.nombres = fila[4] or ""

        jugador.rating = fila[5] or 0

        jugador.nacimiento = fila[6]

        jugador.genero = fila[7] or ""

        jugador.pais = fila[8] or ""

        jugadores.append(jugador)

    return jugadores